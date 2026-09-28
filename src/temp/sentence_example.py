"""Build a sentence-level audit workbook for one statement: dictionary matches next to transformer probabilities."""
import pandas as pd
import yaml
from nltk.stem import PorterStemmer
from nltk.tokenize import word_tokenize
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.auto_econ_sentiment.clean.text_segmentation import TextSegmenter
from src.research_paper.pipes import load_config

COUNTRY, ID = "US", 190
OUT = "src/temp/sentence_audit_US_2020-03-03.xlsx"
config = load_config()
tau = config["transformer"]["sentence_probability_cutoff"]
stemmer = PorterStemmer()
master = yaml.safe_load(open("src/auto_econ_sentiment/data/lexical_master_dict.yaml"))
dicts = [(d, False) for d in config["lexical"]["dictionaries"]["unstemmed"]] + [(d, True) for d in config["lexical"]["dictionaries"]["stemmed"]]
names = {c.split("_")[0]: l for c, l in config["lexical"]["columns"].items()}

def tokens(text, stem):
    toks = [t for t in word_tokenize(text.lower()) if t.isalpha()]
    return [stemmer.stem(t) for t in toks] if stem else toks

def matches(text, name, stem):
    words = master[name]
    pos, neg = set(map(str.lower, map(str, words.get("positive", [])))), set(map(str.lower, map(str, words.get("negative", []))))
    toks = tokens(text, stem)
    return [t for t in toks if t in pos], [t for t in toks if t in neg]

clean = pd.read_parquet(f"data/sentiment/research_paper/2_clean/{COUNTRY}.parquet.gzip")
probs = pd.read_parquet(f"data/sentiment/research_paper/3_sentiment/sentiment_sentences/{COUNTRY}.parquet.gzip")
seg = TextSegmenter(text_column="text_clean", min_chars=config["transformer"]["min_sentence_chars"]).run(clean)
assert len(seg) == len(probs)
seg = pd.concat([seg[["id_text", "sentence_number", "text_clean"]], probs], axis=1)
doc = seg[seg.id_text == ID].reset_index(drop=True)
document = clean[clean.id_text == ID].iloc[0]
panel = pd.read_parquet("data/sentiment/research_paper/panel_sentiment_all.parquet.gzip")
stored = panel[(panel.Country == COUNTRY) & (panel.id_text == ID)].iloc[0]

# Whole-document dictionary counts, to show where sentence sums differ from the document score.
doc_counts = {d: matches(document.text_clean, d, s) for d, s in dicts}
for d, s in dicts:
    p, n = map(len, doc_counts[d])
    recomputed = 1 + (p - n) / (p + n) if p + n else 1
    col = f"{d}_sentiment_posneg" + ("_stem" if s else "")
    assert abs(recomputed - stored[col]) < 1e-9, (d, recomputed, stored[col])

models = [(m["short_name"], m["label"], m["label_map"]) for m in config["transformer"]["models"]]
wb = Workbook()
bold, wrap = Font(bold=True), Alignment(wrap_text=True, vertical="top")
fill = PatternFill("solid", start_color="EEEEEE")

# ---- Sheet 1: sentences
ws = wb.active
ws.title = "Sentences"
header = ["Sentence", "Text"]
for d, _ in dicts:
    header += [f"{names[d]} +words", f"{names[d]} -words", f"{names[d]} P", f"{names[d]} N"]
for short, label, lmap in models:
    header += [f"{label} p({lab}={'+1' if v > 0 else '-1' if v < 0 else '0'})" for lab, v in lmap.items()]
    header += [f"{label} class (p>={tau})", f"{label} d"]
ws.append(header)
for _, row in doc.iterrows():
    out = [int(row.sentence_number), row.text_clean]
    for d, s in dicts:
        p, n = matches(row.text_clean, d, s)
        out += [", ".join(p), ", ".join(n), len(p), len(n)]
    for short, label, lmap in models:
        ps = {lab: float(row[f"{short}_{lab}"]) for lab in lmap}
        out += [round(v, 4) for v in ps.values()]
        cleared = [lab for lab, v in ps.items() if v >= tau]
        out += [cleared[0] if cleared else "none", lmap[cleared[0]] if cleared else 0]
    ws.append(out)
n_rows = len(doc)
first, last = 2, 1 + n_rows
ws.append([])
total_row = last + 2
ws.cell(total_row, 1, "Sum over sentences").font = bold
col_index = {h: i + 1 for i, h in enumerate(header)}
for h, i in col_index.items():
    if h.endswith(" P") or h.endswith(" N"):
        L = get_column_letter(i)
        ws.cell(total_row, i, f"=SUM({L}{first}:{L}{last})").font = bold
for i in range(1, len(header) + 1):
    ws.cell(1, i).font = bold
    ws.cell(1, i).alignment = wrap
    ws.cell(1, i).fill = fill
    ws.column_dimensions[get_column_letter(i)].width = 14
ws.column_dimensions["B"].width = 70
for r in range(2, last + 1):
    ws.cell(r, 2).alignment = wrap
ws.freeze_panes = "C2"

# ---- Sheet 2: document scores with live formulas
ds = wb.create_sheet("Document scores")
ds.append(["Lexical (whole document, as the package scores it)"])
ds.append(["Dictionary", "P (doc)", "N (doc)", "T (tokens)", "PosNeg = (P-N)/(P+N)", "Package PosNeg (stored, 1+...)", "Stored minus 1", "Check",
           "All-Words = (P-N)/T", "Package All-Words minus 1", "Positive words", "Negative words"])
r = 3
from sklearn.feature_extraction.text import CountVectorizer
for d, s in dicts:
    pw, nw = doc_counts[d]
    col = f"{d}_sentiment_posneg" + ("_stem" if s else "")
    aw_col = col.replace("posneg", "allwords")
    text = " ".join(tokens(document.text_clean, s))
    T = int(CountVectorizer(stop_words="english").fit_transform([text]).sum())
    ds.append([names[d], len(pw), len(nw), T,
               f"=IF(B{r}+C{r}=0,0,(B{r}-C{r})/(B{r}+C{r}))", float(stored[col]), f"=F{r}-1", f"=ROUND(E{r}-G{r},9)=0",
               f"=(B{r}-C{r})/D{r}", float(stored[aw_col]) - 1, ", ".join(pw), ", ".join(nw)])
    r += 1
r += 1
ds.cell(r, 1, "Transformers (sentences, tau = %.1f)" % tau).font = bold
r += 1
ds.append(["Model", "N+ ", "N-", "N0", "S (sentences)", "PosNeg = (N+-N-)/(N++N-)", "All-Sentences = (N+-N-)/S", "Stored PosNeg", "Stored All-Sentences", "Check"])
r += 1
sent = "Sentences"
for short, label, lmap in models:
    dcol = get_column_letter(col_index[f"{label} d"])
    ccol = get_column_letter(col_index[f"{label} class (p>={tau})"])
    rng_d, rng_c = f"{sent}!{dcol}{first}:{dcol}{last}", f"{sent}!{ccol}{first}:{ccol}{last}"
    ds.append([label, f'=COUNTIF({rng_d},1)', f'=COUNTIF({rng_d},-1)', f'=COUNTIFS({rng_d},0,{rng_c},"<>none")', n_rows,
               f"=IF(B{r}+C{r}=0,\"undefined\",(B{r}-C{r})/(B{r}+C{r}))", f"=(B{r}-C{r})/E{r}",
               None if pd.isna(stored[f"{short}_net_sentiment_posneg"]) else float(stored[f"{short}_net_sentiment_posneg"]),
               float(stored[f"{short}_net_sentiment_allsentences"]), f"=ROUND(G{r}-I{r},9)=0"])
    r += 1
for row in ds.iter_rows(min_row=1, max_row=r):
    for c in row:
        if c.row in (1, 2) or (isinstance(c.value, str) and c.value in ("Model", "Dictionary")):
            c.font = bold
for i in range(1, 13):
    ds.column_dimensions[get_column_letter(i)].width = 18

# ---- Sheet 3: notes
notes = wb.create_sheet("Notes")
for line in [
    f"Statement: FOMC, {document.date.date()} (id_text {ID}), {n_rows} sentences of at least {config['transformer']['min_sentence_chars']} characters.",
    "Sentences come from re-segmenting 2_clean text; transformer probabilities are the cached 51-bank run, aligned row for row.",
    "Dictionary matches use the package tokenizer (lower-case, alphabetic tokens); Bennani-Neuenkirch and Apel-Blix Grimaldi match Porter stems.",
    "The document score counts matches on the whole text; the sentence sums can differ slightly because sentences under the length floor are dropped.",
    "A sentence counts toward a class only if that class probability >= tau; column 'd' is the class direction (+1, 0, -1), 0 when nothing clears tau.",
    "FOMC-RoBERTa: LABEL_0 dovish (-1), LABEL_1 hawkish (+1), LABEL_2 neutral (0).",
    "The 'Check' columns compare the Excel formula with the value stored in panel_sentiment_all.parquet.gzip.",
]:
    notes.append([line])
notes.column_dimensions["A"].width = 150
wb.save(OUT)
print(OUT)
print(doc[["sentence_number", "text_clean"]].to_string())
